from __future__ import annotations

import base64
from pathlib import Path
from typing import Any


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in {'.txt', '.md', '.json'}:
        return path.read_text(encoding='utf-8')
    if suffix in {'.csv', '.tsv', '.xlsx'}:
        import pandas as pd
        if suffix == '.xlsx':
            sheets = pd.read_excel(path, sheet_name=None, nrows=100)
        else:
            sheets = {path.name: pd.read_csv(path, sep='\t' if suffix == '.tsv' else ',', nrows=100)}
        return '\n\n'.join(f'Sheet: {name}\n{frame.to_string(index=False)}' +
                           ('\n[Preview limited to 100 rows.]' if len(frame) == 100 else '')
                           for name, frame in sheets.items())
    if suffix == '.pdf':
        from pypdf import PdfReader
        return '\n'.join(page.extract_text() or '' for page in PdfReader(path).pages).strip() or '无可提取文字'
    if suffix == '.docx':
        from docx import Document
        document = Document(path)
        return '\n'.join([paragraph.text for paragraph in document.paragraphs if paragraph.text] +
                         ['\n'.join(' | '.join(cell.text for cell in row.cells) for row in table.rows)
                          for table in document.tables])
    return 'Opaque attachment; read the referenced file with a supported tool. Its contents have not been extracted.'


def content_blocks(sources: list[tuple[dict[str, str], Path]], text: str, workspace_root: Path) -> str | list[dict[str, Any]]:
    if not sources:
        return text
    blocks: list[dict[str, Any]] = [{'type': 'text', 'text': text}]
    text_items = []
    marker = '\n[Attachment context truncated at 20,000 characters.]'
    for ref, path in sources:
        if path.suffix.lower() in {'.png', '.jpg', '.jpeg', '.webp', '.gif'}:
            data = base64.b64encode(path.read_bytes()).decode('ascii')
            blocks.append({'type': 'image_url', 'image_url': {'url': f'data:{ref["media_type"]};base64,{data}'}})
            continue
        header = f'Attachment: {ref["name"]}\nPath: /{path.relative_to(workspace_root).as_posix()}\n'
        text_items.append((len(blocks), header, extract_text(path)))
        blocks.append({'type': 'text', 'text': ''})
    body_budget = 20000 - sum(len(header) for _, header, _ in text_items)
    truncated = sum(len(body) for _, _, body in text_items) > body_budget
    remaining = max(0, body_budget - (len(marker) if truncated else 0))
    for index, header, body in text_items:
        excerpt = body[:remaining]
        remaining -= len(excerpt)
        blocks[index]['text'] = header + excerpt
    if truncated:
        blocks[text_items[-1][0]]['text'] += marker
    return blocks
